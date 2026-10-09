"""Import and reconcile the partner's private Access outputs stage by stage."""
from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from .reference_onezero import ReferenceOneZeroRun
from .reference_step1 import ReferenceStep1Run


def _normalize(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


_HEADER_ALIASES = {
    "employeeid": "Employee_ID",
    "tccemployeeid": "Employee_ID",
    "totalhours": "TotalHours",
    "totalwages": "TotalWages",
    "firststart": "FirstStart",
    "lastend": "LastEnd",
    "days": "Days",
    "weeks": "DenominatorWeeks",
    "weeksworked": "DenominatorWeeks",
    "weeksworks": "DenominatorWeeks",
    "avgweeklyhours": "AvgWeeklyHours",
    "averageweeklyhours": "AvgWeeklyHours",
    "avgweeklywage": "AvgWeeklyWages",
    "avgweeklywages": "AvgWeeklyWages",
    "averageweeklywage": "AvgWeeklyWages",
    "averageweeklywages": "AvgWeeklyWages",
    "hourstestpass": "PassedHours",
    "passedhours": "PassedHours",
    "wagetestpass": "PassedWages",
    "passedwages": "PassedWages",
    "overalleligible": "Eligible",
    "eligible": "Eligible",
    "paycycle": "PayCycle",
    "payfrequency": "PayCycle",
    "includedperiodcount": "IncludedPeriodCount",
    "majorperiodcount": "IncludedPeriodCount",
    "unit": "Unit_ID",
    "unitid": "Unit_ID",
    "employerlocationid": "Unit_ID",
    "tccunitid": "Unit_ID",
    "fein": "FEIN",
    "total": "Total",
    "average": "Average",
    "averageheadcount": "Average",
}
for index in range(1, 13):
    _HEADER_ALIASES[f"month{index}"] = f"Month{index}"


_STAGE_ALIASES = {
    "commonemployeetotals": "common_employee_totals",
    "totalhoursandwages": "common_employee_totals",
    "minmaxeligibility": "min_max_eligibility",
    "minmaxeligibleemployees": "min_max_eligibility",
    "payperiodeligibility": "pay_period_eligibility",
    "payperiodeligibleemployees": "pay_period_eligibility",
    "minmaxonezerosummary": "min_max_onezero_summary",
    "minmaxsummary": "min_max_onezero_summary",
    "payperiodonezerosummary": "pay_period_onezero_summary",
    "payperiodsummary": "pay_period_onezero_summary",
}


@dataclass(frozen=True)
class StageSpec:
    name: str
    keys: tuple[str, ...]
    fields: tuple[str, ...]
    numeric_fields: tuple[str, ...] = ()
    boolean_fields: tuple[str, ...] = ()
    tolerance: Decimal = Decimal("0")


_STAGE_SPECS = (
    StageSpec(
        "common_employee_totals",
        ("Employee_ID",),
        ("TotalHours", "TotalWages"),
        numeric_fields=("TotalHours", "TotalWages"),
        tolerance=Decimal("0.01"),
    ),
    StageSpec(
        "min_max_eligibility",
        ("Employee_ID",),
        (
            "TotalHours", "TotalWages", "DenominatorWeeks", "AvgWeeklyHours",
            "AvgWeeklyWages", "PassedHours", "PassedWages", "Eligible",
        ),
        numeric_fields=(
            "TotalHours", "TotalWages", "DenominatorWeeks", "AvgWeeklyHours",
            "AvgWeeklyWages",
        ),
        boolean_fields=("PassedHours", "PassedWages", "Eligible"),
        tolerance=Decimal("0.01"),
    ),
    StageSpec(
        "pay_period_eligibility",
        ("Employee_ID",),
        (
            "TotalHours", "TotalWages", "DenominatorWeeks", "AvgWeeklyHours",
            "AvgWeeklyWages", "PassedHours", "PassedWages", "Eligible",
        ),
        numeric_fields=(
            "TotalHours", "TotalWages", "DenominatorWeeks", "AvgWeeklyHours",
            "AvgWeeklyWages",
        ),
        boolean_fields=("PassedHours", "PassedWages", "Eligible"),
        tolerance=Decimal("0.01"),
    ),
    StageSpec(
        "min_max_onezero_summary",
        ("Unit_ID", "FEIN"),
        tuple([f"Month{index}" for index in range(1, 13)] + ["Total", "Average"]),
        numeric_fields=tuple([f"Month{index}" for index in range(1, 13)] + ["Total", "Average"]),
        tolerance=Decimal("0.01"),
    ),
    StageSpec(
        "pay_period_onezero_summary",
        ("Unit_ID", "FEIN"),
        tuple([f"Month{index}" for index in range(1, 13)] + ["Total", "Average"]),
        numeric_fields=tuple([f"Month{index}" for index in range(1, 13)] + ["Total", "Average"]),
        tolerance=Decimal("0.01"),
    ),
)


def _canonical_header(value: Any) -> str:
    normalized = _normalize(value)
    return _HEADER_ALIASES.get(normalized, str(value or "").strip())


def _canonical_stage(value: str) -> str | None:
    normalized = _normalize(Path(value).stem)
    return _STAGE_ALIASES.get(normalized, normalized if normalized in _STAGE_ALIASES.values() else None)


def _clean_rows(headers: list[Any], values: Iterable[Iterable[Any]]) -> list[dict[str, Any]]:
    canonical = [_canonical_header(header) for header in headers]
    rows: list[dict[str, Any]] = []
    for value_row in values:
        row = {canonical[index]: value for index, value in enumerate(value_row) if index < len(canonical)}
        if any(value not in (None, "") for value in row.values()):
            rows.append(row)
    return rows


def _read_csv(path: Path) -> list[dict[str, Any]]:
    with path.open(newline="", encoding="utf-8-sig") as source:
        reader = csv.reader(source)
        headers = next(reader, [])
        return _clean_rows(headers, reader)


def _read_xlsx(path: Path) -> dict[str, list[dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    stages: dict[str, list[dict[str, Any]]] = {}
    try:
        for sheet in workbook.worksheets:
            stage = _canonical_stage(sheet.title)
            if stage is None:
                continue
            iterator = sheet.iter_rows(values_only=True)
            headers = list(next(iterator, ()))
            stages[stage] = _clean_rows(headers, iterator)
    finally:
        workbook.close()
    return stages


def load_access_bundle(path: Path) -> dict[str, list[dict[str, Any]]]:
    """Load a directory of canonical CSVs or one multi-sheet XLSX workbook."""
    path = Path(path)
    if path.is_file() and path.suffix.lower() in {".xlsx", ".xlsm"}:
        stages = _read_xlsx(path)
    elif path.is_dir():
        stages = {}
        for file in sorted(path.glob("*.csv")):
            stage = _canonical_stage(file.name)
            if stage:
                stages[stage] = _read_csv(file)
        for file in sorted(path.glob("*.xlsx")):
            stages.update(_read_xlsx(file))
    else:
        raise FileNotFoundError("Access output path must be a directory or XLSX workbook")
    if not stages:
        raise ValueError("no recognized Access output stages found")
    return stages


def _decimal(value: Any) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    text = str(value).strip().replace("$", "").replace(",", "")
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError("invalid numeric benchmark value") from exc


def _boolean(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    normalized = _normalize(value)
    if normalized in {"1", "true", "yes", "y", "eligible", "pass", "passed"}:
        return True
    if normalized in {"0", "false", "no", "n", "ineligible", "fail", "failed", ""}:
        return False
    raise ValueError("invalid boolean benchmark value")


def _key(row: dict[str, Any], fields: tuple[str, ...]) -> tuple[str, ...]:
    return tuple(str(row.get(field, "")).strip() for field in fields)


def _computed_stages(step1: ReferenceStep1Run, onezero: ReferenceOneZeroRun) -> dict[str, list[dict[str, Any]]]:
    totals = [
        {
            "Employee_ID": employee_id,
            "TotalHours": str(row.total_hours),
            "TotalWages": str(row.total_wages),
        }
        for employee_id, row in sorted(step1.common_totals.items())
    ]

    def eligibility(rows: dict[str, Any]) -> list[dict[str, Any]]:
        return [
            {
                "Employee_ID": employee_id,
                "TotalHours": str(row.total_hours),
                "TotalWages": str(row.total_wages),
                "DenominatorWeeks": str(row.denominator_weeks),
                "AvgWeeklyHours": str(row.rounded_avg_weekly_hours),
                "AvgWeeklyWages": str(row.rounded_avg_weekly_wages),
                "PassedHours": row.passed_hours,
                "PassedWages": row.passed_wages,
                "Eligible": row.eligible,
            }
            for employee_id, row in sorted(rows.items())
        ]

    def summary(result: Any) -> list[dict[str, Any]]:
        output: list[dict[str, Any]] = []
        for row in result.summary:
            payload = {
                "Unit_ID": row.unit_id,
                "FEIN": row.fein,
                "Total": row.total,
                "Average": str(row.application_average),
            }
            payload.update({f"Month{index + 1}": value for index, value in enumerate(row.month_totals)})
            output.append(payload)
        return output

    return {
        "common_employee_totals": totals,
        "min_max_eligibility": eligibility(step1.min_max),
        "pay_period_eligibility": eligibility(step1.pay_period),
        "min_max_onezero_summary": summary(onezero.min_max),
        "pay_period_onezero_summary": summary(onezero.pay_period),
    }


def compare_stage(
    benchmark: list[dict[str, Any]],
    computed: list[dict[str, Any]],
    spec: StageSpec,
) -> dict[str, Any]:
    benchmark_by_key = {_key(row, spec.keys): row for row in benchmark}
    computed_by_key = {_key(row, spec.keys): row for row in computed}
    differences: list[dict[str, Any]] = []
    exact_rows = 0
    for row_key in sorted(set(benchmark_by_key) | set(computed_by_key)):
        expected = benchmark_by_key.get(row_key)
        actual = computed_by_key.get(row_key)
        if expected is None or actual is None:
            differences.append({
                "key": list(row_key),
                "field": None,
                "benchmark": expected,
                "computed": actual,
                "status": "missing_benchmark" if expected is None else "missing_computed",
            })
            continue
        row_differences = 0
        for field in spec.fields:
            if field in spec.numeric_fields:
                expected_value = _decimal(expected.get(field))
                actual_value = _decimal(actual.get(field))
                difference = actual_value - expected_value
                if abs(difference) > spec.tolerance:
                    differences.append({
                        "key": list(row_key),
                        "field": field,
                        "benchmark": str(expected_value),
                        "computed": str(actual_value),
                        "difference": str(difference),
                        "status": "mismatch",
                    })
                    row_differences += 1
            elif field in spec.boolean_fields:
                expected_value = _boolean(expected.get(field))
                actual_value = _boolean(actual.get(field))
                if expected_value != actual_value:
                    differences.append({
                        "key": list(row_key),
                        "field": field,
                        "benchmark": expected_value,
                        "computed": actual_value,
                        "status": "mismatch",
                    })
                    row_differences += 1
            elif str(expected.get(field, "")).strip() != str(actual.get(field, "")).strip():
                differences.append({
                    "key": list(row_key),
                    "field": field,
                    "benchmark": expected.get(field),
                    "computed": actual.get(field),
                    "status": "mismatch",
                })
                row_differences += 1
        if row_differences == 0:
            exact_rows += 1
    return {
        "stage": spec.name,
        "keys": list(spec.keys),
        "fields": list(spec.fields),
        "tolerance": str(spec.tolerance),
        "benchmark_rows": len(benchmark),
        "computed_rows": len(computed),
        "exact_rows": exact_rows,
        "difference_count": len(differences),
        "reconciled": not differences,
        "differences": differences,
    }


def compare_access_bundle(
    access_stages: dict[str, list[dict[str, Any]]],
    step1: ReferenceStep1Run,
    onezero: ReferenceOneZeroRun,
) -> dict[str, Any]:
    computed = _computed_stages(step1, onezero)
    results: list[dict[str, Any]] = []
    first_divergent_stage: str | None = None
    for spec in _STAGE_SPECS:
        if spec.name not in access_stages:
            continue
        result = compare_stage(access_stages[spec.name], computed[spec.name], spec)
        results.append(result)
        if not result["reconciled"] and first_divergent_stage is None:
            first_divergent_stage = spec.name
    return {
        "step1_calculation_id": step1.calculation_id,
        "onezero_calculation_id": onezero.calculation_id,
        "stages_compared": [result["stage"] for result in results],
        "first_divergent_stage": first_divergent_stage,
        "reconciled": bool(results) and all(result["reconciled"] for result in results),
        "stage_summaries": [
            {key: value for key, value in result.items() if key != "differences"}
            for result in results
        ],
        "stages": results,
    }
