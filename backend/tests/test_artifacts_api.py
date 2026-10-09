"""Claim packages and governed downloads. [BLOB-3, BLOB-4, D16, D24]

Endpoint tests run the real app with a local byte store, a SQLite repository
and a fake Gold source; the workbook test opens what was actually generated.
"""
from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from application import wiring
from application.artifacts import ArtifactService
from application.claim_package import SHEETS, build_claim_package
from application.identity import AllClientsAuthorizer, AllowListAuthorizer, ANONYMOUS
from infrastructure.persistence.sqlite_artifacts import SqliteArtifactRepository
from infrastructure.storage.local_store import LocalArtifactStore

RUN = {
    "run_id": "0123456789abcdef0123456789abcdef", "program_code": "GA_JTC", "tax_year": 2024,
    "fiscal_year_begin": "2024-01-01", "fiscal_year_end": "2024-12-31",
    "rule_pack_version": "rules/ga_jtc/2024.json", "step1_assumption_set_id": "GDS-STEP1-v8",
    "onezero_assumption_set_id": "GDS-ONEZERO-v7", "code_commit": "a" * 40, "calculation_id": "calc-1",
    "run_started_at": datetime(2026, 10, 1, 9, 0), "run_completed_at": datetime(2026, 10, 1, 9, 5),
}

#: The synthetic payroll's employee-id shape, plus anything SSN- or email-shaped.
EMPLOYEE_ID = re.compile(r"^E\d{5}$")
SSN = re.compile(r"\b\d{3}-\d{2}-\d{4}\b")
EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")
IDENTIFIER_HEADER = re.compile(r"employee|ssn|social|first_name|last_name|record_token|email|dob", re.I)


class FakeSource:
    """Gold-shaped rows, deliberately carrying fields that must NOT reach the file."""

    def __init__(self, run=RUN):
        self.run = run

    def latest_run(self, client_id):
        return self.run

    def coverage(self, client_id, tax_year):
        return [{
            "program_code": "GA_JTC", "tax_year": tax_year, "status": "assessed", "provenance": "computed",
            "provenance_verified": True, "identified_usd": "31500", "claimed_usd": "0", "certified_usd": "0",
            "outcome": "unclaimed", "is_base_year": False, "decided_by": "someone@example.com",
        }]

    def locations(self, client_id, tax_year):
        return [{
            "method": m, "worksite_id": "GA_SAV_DC", "program_code": "GA_JTC", "tax_year": tax_year,
            "status": "above_minimum", "average_headcount": 61.5, "eligible_jobs": 9.0,
            "per_job_usd": Decimal("3500"), "annual_credit_usd": Decimal("31500"), "run_id": RUN["run_id"],
            "employee_id": "E00200", "record_token": "deadbeef",   # must be dropped
        } for m in ("min_max", "pay_period")]

    def installments(self, client_id, tax_year):
        return [{
            "method": "min_max", "worksite_id": "GA_SAV_DC", "program_code": "GA_JTC", "creation_year": 2023,
            "tax_year": tax_year, "tranche_jobs": 9.0, "jobs_maintained": 9.0, "per_job_usd": Decimal("3500"),
            "claimable_usd": Decimal("31500"), "projected": False, "run_id": RUN["run_id"],
        }]


@pytest.fixture
def service(tmp_path):
    repo = SqliteArtifactRepository(tmp_path / "artifacts.db")
    svc = ArtifactService(LocalArtifactStore(tmp_path / "blobs"), repo, AllClientsAuthorizer())
    wiring.set_artifact_service(svc)
    wiring.set_claim_package_source(FakeSource())
    yield svc
    wiring.set_artifact_service(None)
    wiring.set_claim_package_source(None)
    wiring.set_client_authorizer(None)
    repo.close()


@pytest.fixture
def api(monkeypatch):
    import main

    monkeypatch.delenv("AUTH_MODE", raising=False)
    return TestClient(main.app)


def _create(api, client_id=1001, year=2024):
    response = api.post(f"/api/clients/{client_id}/packages", params={"tax_year": year})
    assert response.status_code == 201, response.text
    return response.json()


# ----------------------------------------------------------------- endpoints


def test_create_then_download_serves_the_workbook_and_records_it(api, service):
    created = _create(api)
    assert created["kind"] == "claim_package" and created["client_id"] == 1001 and created["tax_year"] == 2024

    response = api.get(f"/api/clients/1001/artifacts/{created['artifact_id']}")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in response.headers["content-disposition"]
    assert response.headers["x-content-sha256"] == created["sha256"]
    load_workbook(io.BytesIO(response.content))  # a valid XLSX
    assert [a.outcome for a in service.repository.accesses(1001)] == ["served"]


def test_the_blob_lives_under_the_clients_prefix(api, service):
    created = _create(api)
    artifact = service.repository.find(created["artifact_id"])
    assert artifact.blob_path.startswith(f"1001/{RUN['run_id']}/")
    assert artifact.container_name == "outputs" and artifact.retain_until is not None


def test_an_unauthorized_caller_is_denied_and_it_is_recorded(api, service):
    created = _create(api)
    service.authorizer = AllowListAuthorizer({})  # the fake authorizer (D24)
    response = api.get(f"/api/clients/1001/artifacts/{created['artifact_id']}")
    assert response.status_code == 403
    assert [a.outcome for a in service.repository.accesses(1001)] == ["denied"]


def test_another_clients_path_is_denied_even_when_both_are_authorized(api, service):
    created = _create(api, client_id=1001)
    response = api.get(f"/api/clients/2001/artifacts/{created['artifact_id']}")
    assert response.status_code == 403
    accesses = service.repository.accesses(1001)
    assert [a.outcome for a in accesses] == ["denied"], "recorded against the artifact's owner"


def test_a_registered_artifact_whose_bytes_are_gone_is_missing(api, service, tmp_path):
    created = _create(api)
    artifact = service.repository.find(created["artifact_id"])
    (tmp_path / "blobs" / "outputs" / artifact.blob_path).unlink()
    response = api.get(f"/api/clients/1001/artifacts/{created['artifact_id']}")
    assert response.status_code == 404
    assert [a.outcome for a in service.repository.accesses(1001)] == ["missing"]


def test_an_unregistered_id_is_404_and_writes_no_access_row(api, service):
    assert api.get("/api/clients/1001/artifacts/nope").status_code == 404
    assert service.repository.accesses(1001) == []


def test_a_blob_with_no_row_cannot_be_downloaded(api, service):
    service.store.put("outputs", "1001/run/orphan.xlsx", b"x", "application/octet-stream")
    assert api.get("/api/clients/1001/artifacts/orphan.xlsx").status_code == 404


def test_a_client_never_calculated_gets_409_not_an_empty_workbook(api, service):
    wiring.set_claim_package_source(FakeSource(run=None))
    response = api.post("/api/clients/1001/packages", params={"tax_year": 2024})
    assert response.status_code == 409


def test_a_second_package_supersedes_the_first(api, service):
    first = _create(api)
    second = _create(api)
    live = service.repository.live(1001, kind="claim_package", tax_year=2024)
    assert [a.artifact_id for a in live] == [second["artifact_id"]]
    assert service.repository.find(first["artifact_id"]).superseded_at is not None


# ----------------------------------------------------------------- the workbook


def _workbook():
    package = build_claim_package(1001, 2024, FakeSource(), generated_by="analyst@example.com",
                                  generated_at=datetime(2026, 10, 8, tzinfo=timezone.utc))
    return package, load_workbook(io.BytesIO(package.data))


def test_the_workbook_has_exactly_the_five_sheets():
    package, book = _workbook()
    assert book.sheetnames == list(SHEETS)
    assert package.rows["Locations"] == 2 and package.rows["Installments"] == 1 and package.rows["Coverage"] == 1


def test_provenance_names_the_run_rules_and_assumptions():
    _, book = _workbook()
    fields = {row[0]: row[1] for row in book["Provenance"].iter_rows(min_row=2, values_only=True)}
    assert fields["run_id"] == RUN["run_id"]
    assert fields["rule_pack_version"] == "rules/ga_jtc/2024.json"
    assert fields["step1_assumption_set_id"] == "GDS-STEP1-v8"
    assert fields["generated_by"] == "analyst@example.com"
    assert "trail_sha256" in fields


def test_no_employee_identifier_reaches_the_workbook():
    """Unit grain only (D16): no identifier column, and no identifier-shaped value."""
    _, book = _workbook()
    for sheet in book.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if sheet.title in ("Coverage", "Locations", "Installments"):
            headers = [h for h in rows[0] if h]
            assert not [h for h in headers if IDENTIFIER_HEADER.search(str(h))], (sheet.title, headers)
        for row in rows:
            for value in row:
                text = "" if value is None else str(value)
                assert not EMPLOYEE_ID.match(text), (sheet.title, text)
                assert not SSN.search(text), (sheet.title, text)
                if sheet.title != "Provenance":  # generated_by is the analyst, not an employee
                    assert not EMAIL.search(text), (sheet.title, text)


def test_money_lands_as_numbers():
    _, book = _workbook()
    header = [c.value for c in book["Locations"][1]]
    value = book["Locations"].cell(row=2, column=header.index("annual_credit_usd") + 1).value
    assert value == 31500


def test_a_denied_request_never_reads_the_bytes(tmp_path):
    from application.artifacts import Artifact

    class ExplodingStore(LocalArtifactStore):
        def get(self, container, path):
            raise AssertionError("bytes must not be read for a denied request")

    repo = SqliteArtifactRepository(tmp_path / "a.db")
    svc = ArtifactService(ExplodingStore(tmp_path / "b"), repo, AllClientsAuthorizer())
    artifact = repo.record(Artifact(
        kind="claim_package", client_id=1001, run_id="r", container_name="outputs",
        blob_path="1001/r/p.xlsx", file_name="p.xlsx", content_type="t", size_bytes=1,
        sha256="a" * 64, created_by="x",
    ))
    assert svc.get(ANONYMOUS, 2001, artifact.artifact_id).outcome == "denied"
    repo.close()
