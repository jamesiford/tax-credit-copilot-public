"""Claim package v1: one XLSX per client x tax year, unit grain only. [BLOB-4, D16]

Five sheets, read from what the run already wrote — no engine runs here:

- **Coverage**     the client's program claims for the year (`tcc_ops.program_claim`)
                   with the derived outcome, exactly as the coverage matrix reads them.
- **Locations**    the current unit-year rows (`tcc_gold.unit_year_master`), both methods.
- **Installments** the current credit tranches paying out in that year
                   (`tcc_gold.unit_credit_installment`).
- **Provenance**   which run, rule pack and assumption sets produced every figure.
- **Disclaimer**   what the workbook is and is not.

**No employee identifiers, names or tax ids ever enter the workbook.** Columns
are taken from an explicit allow-list per sheet rather than "whatever the row
holds", so a column added to Gold later cannot leak into a file a client keeps.
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal
from typing import Any, Iterable, Protocol

SHEETS = ("Coverage", "Locations", "Installments", "Provenance", "Disclaimer")

COVERAGE_COLUMNS = (
    "program_code", "tax_year", "status", "provenance", "provenance_verified",
    "identified_usd", "claimed_usd", "certified_usd", "outcome", "is_base_year",
    "average_headcount", "base_average_headcount", "threshold_jobs",
    "identified_run_id", "rule_pack_version",
)
LOCATION_COLUMNS = (
    "method", "worksite_id", "program_code", "tax_year", "status", "average_headcount",
    "is_base_year", "base_average_headcount", "base_year_source", "net_new_vs_base",
    "months_reachable", "year_complete", "threshold", "threshold_source", "qualified",
    "tier", "designation", "tier_source", "eligible_jobs", "new_jobs_created",
    "high_water_mark", "per_job_usd", "tranche_years", "annual_credit_usd", "run_id",
)
INSTALLMENT_COLUMNS = (
    "method", "worksite_id", "program_code", "creation_year", "tax_year", "tranche_jobs",
    "jobs_maintained", "per_job_usd", "claimable_usd", "projected", "run_id",
)
PROVENANCE_KEYS = (
    "run_id", "program_code", "tax_year", "fiscal_year_begin", "fiscal_year_end",
    "rule_pack_version", "step1_assumption_set_id", "onezero_assumption_set_id",
    "code_commit", "calculation_id", "run_started_at", "run_completed_at",
)

DISCLAIMER = (
    "Tax Credit Copilot claim package (v1).",
    "Every figure was computed by the deterministic calculation engine from payroll, "
    "under the rule pack and assumption sets named on the Provenance sheet. No figure "
    "was produced by a language model.",
    "Claimed amounts marked client_asserted or inherited rest on an unverified assertion; "
    "any variance computed against them is only as reliable as that assertion.",
    "The workbook is at location grain. It contains no employee identifiers, names or tax ids.",
    "This is not tax advice. A return must be prepared and reviewed by a qualified professional.",
)


class ClaimPackageSource(Protocol):
    def latest_run(self, client_id: int) -> dict[str, Any] | None: ...

    def coverage(self, client_id: int, tax_year: int) -> list[dict[str, Any]]: ...

    def locations(self, client_id: int, tax_year: int) -> list[dict[str, Any]]: ...

    def installments(self, client_id: int, tax_year: int) -> list[dict[str, Any]]: ...


class NoRunForClient(LookupError):
    """The client has never been calculated: there is nothing to package."""


@dataclass(frozen=True)
class ClaimPackage:
    data: bytes
    file_name: str
    run_id: str
    program_code: str | None
    rows: dict[str, int]


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _cell(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value) if value == value.to_integral_value() or abs(value) < Decimal("1e15") else str(value)
    if isinstance(value, datetime):
        return value.astimezone(timezone.utc).replace(tzinfo=None) if value.tzinfo else value
    if isinstance(value, (str, int, float, bool, date)) or value is None:
        return value
    return str(value)


def _table(sheet: Any, columns: Iterable[str], rows: Iterable[dict[str, Any]]) -> int:
    columns = list(columns)
    sheet.append(columns)
    count = 0
    for row in rows:
        sheet.append([_cell(row.get(c)) for c in columns])
        count += 1
    sheet.freeze_panes = "A2"
    return count


def build_claim_package(
    client_id: int,
    tax_year: int,
    source: ClaimPackageSource,
    *,
    generated_by: str,
    generated_at: datetime | None = None,
) -> ClaimPackage:
    from openpyxl import Workbook

    run = source.latest_run(client_id)
    if run is None:
        raise NoRunForClient(f"client {client_id} has no completed calculation run to package")
    generated_at = generated_at or datetime.now(timezone.utc)

    book = Workbook()
    coverage = book.active
    coverage.title = "Coverage"
    rows = {"Coverage": _table(coverage, COVERAGE_COLUMNS, source.coverage(client_id, tax_year))}
    rows["Locations"] = _table(book.create_sheet("Locations"), LOCATION_COLUMNS, source.locations(client_id, tax_year))
    rows["Installments"] = _table(
        book.create_sheet("Installments"), INSTALLMENT_COLUMNS, source.installments(client_id, tax_year)
    )

    provenance = book.create_sheet("Provenance")
    provenance.append(["field", "value"])
    provenance.append(["client_id", client_id])
    provenance.append(["package_tax_year", tax_year])
    for key in PROVENANCE_KEYS:
        provenance.append([key, _cell(run.get(key))])
    provenance.append(["trail_sha256", _cell(run.get("trail_sha256")) or "not recorded on the run manifest"])
    provenance.append(["generated_by", generated_by])
    provenance.append(["generated_at", _cell(generated_at)])
    provenance.append(["package_format", "claim_package v1"])
    rows["Provenance"] = provenance.max_row - 1

    disclaimer = book.create_sheet("Disclaimer")
    for line in DISCLAIMER:
        disclaimer.append([line])
    rows["Disclaimer"] = len(DISCLAIMER)

    for sheet in book.worksheets:
        sheet.column_dimensions["A"].width = 28

    buffer = io.BytesIO()
    book.save(buffer)
    run_id = str(run["run_id"])
    return ClaimPackage(
        data=buffer.getvalue(),
        file_name=f"claim-package-{client_id}-TY{tax_year}-{run_id[:8]}.xlsx",
        run_id=run_id,
        program_code=run.get("program_code"),
        rows=rows,
    )


class GoldClaimPackageSource:
    """Reads the package from Gold and the claim store. No engine, no payroll."""

    METHODS = ("min_max", "pay_period")

    def __init__(self, reader: Any, claims: Any) -> None:
        self.reader = reader
        self.claims = claims

    def latest_run(self, client_id: int) -> dict[str, Any] | None:
        return self.reader.latest_run(client_id)

    def coverage(self, client_id: int, tax_year: int) -> list[dict[str, Any]]:
        return [c.to_dict() for c in self.claims.coverage(client_id, tax_years=[tax_year])]

    def locations(self, client_id: int, tax_year: int) -> list[dict[str, Any]]:
        out: list[dict[str, Any]] = []
        for method in self.METHODS:
            items, _ = self.reader.unit_years(client_id, method=method, tax_year=tax_year, limit=100_000)
            out.extend({**row, "method": method} for row in items)
        return out

    def installments(self, client_id: int, tax_year: int) -> list[dict[str, Any]]:
        out: list[dict[str, Any]] = []
        for method in self.METHODS:
            out.extend(
                {**row, "method": method}
                for row in self.reader.installments(client_id, method=method)
                if int(row.get("tax_year") or 0) == tax_year
            )
        return out
